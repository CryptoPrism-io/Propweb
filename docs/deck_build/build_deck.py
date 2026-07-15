# -*- coding: utf-8 -*-
"""Rebuilds docs/PropWeb_Pitch_v2.pptx from the pristine 9-slide source,
adding the new Act 2 (+Market) slides and presenter notes per
docs/superpowers/specs/2026-07-14-pitch-deck-design.md."""

from pptx import Presentation
from deck_helpers import add_blank_slide, add_header, add_stat_row, add_mixed_textbox, add_card, add_textbox, set_notes, get_slide_text, add_table, remove_slide_by_text, add_flow_steps, reorder_slides

SOURCE = 'source.pptx'
OUTPUT = '../PropWeb_Pitch_v2.pptx'

FINAL_ORDER_MARKERS = [
    'PROPWEB',
    'Renting in India is broken',
    'The hidden cost is trust',
    'A large market with an unsolved trust gap',
    'A trust-first rental platform with an AI engine',
    'From verified profile to matched move-in',
    'Problems no one in India solves well',
    'Two engines: connect today, services tomorrow',
    'One well-organised building, not eleven services',
    'A modern stack, corrected for India economics',
    'Verification, the legal way',
    'Four phases, one city first',
    'Small team, sequenced hiring',
    'Three scenarios, months 0–12',
    'What we need from you, next 90 days',
    'Existing portals sell leads',
]

EXISTING_NOTES = {
    'PROPWEB': "Open cold: this is the one-line pitch. Say it, don't read it — \"PropWeb is India's trust-first, AI-powered rental platform: verified owners meet verified tenants directly, no brokers, no fake listings, no spam calls.\" Pause after the tagline before moving to the problem slide — let the promise land before you justify it.",
    'Renting in India is broken': "Walk both columns — don't just read the tenant side. The room already knows renting is broken; the point is naming both sides of the pain so the solution slide lands as a genuine two-sided fix, not a tenant-only app. If asked \"isn't this just NoBroker,\" hold that for the next slide.",
    'The hidden cost is trust': "These four numbers carry the slide — let them breathe, don't rush to the NoBroker line. The NoBroker callout is the pre-emptive answer to \"hasn't this already been solved\" — say revenue and loss together, in the same breath: proof of demand, proof the trust layer is still broken.",
    'A trust-first rental platform with an AI engine': "Six pillars, but only two are truly new to the room: Pay-to-Connect and the Trust Score — spend your time there. Verified badges and matchmaking are easy to nod along to; the pay-to-connect economics are what actually kills fake enquiries, so be ready to defend the fee-amount question.",
    'From verified profile to matched move-in': "This is the natural hand-off to the live demo — after \"Connect,\" say \"and I can show you this right now\" and switch to the demo app. Have it already loaded to the Koramangala search before this slide so the transition has no dead air.",
    'Problems no one in India solves well': "This slide is doing double duty as the moat argument, along with the closing slide — scam prevention, background checks and digital agreements are what a trust-first platform can sell that a lead-selling portal structurally can't retrofit. If a CEO asks \"why can't NoBroker just copy this,\" point back here.",
    'Two engines: connect today, services tomorrow': "Two engines, two timeframes — be explicit that \"today\" funds the company and \"tomorrow\" is where the real upside is, using NoBroker's own ~50% non-listing mix as the proof point. Don't let this slide turn into a pricing negotiation; final pricing is explicitly TBD with the CEOs.",
    'Existing portals sell leads': "This is the mic-drop — deliver the two lines slowly, then stop talking. Let \"Existing portals sell leads. PropWeb sells trust.\" sit before you open the floor. If the ask & close slide raised a decision, circle back to it once more before Q&A.",
}

MARKET_NOTES = ("This slide exists to pre-empt the \"is this market even big enough\" question before it's asked — "
                "lead with the CAGR, not the dollar figure; growth rate is the more persuasive number in the room. "
                "Bengaluru isn't arbitrary: name the three reasons (rental velocity, brokerage pain, digital readiness) "
                "so it reads as a decision, not a default.")


ARCHITECTURE_NOTES = ("Keep this slide short on stage — the point is reassurance, not a systems-design lecture. "
                       "One sentence: \"one well-built application, not eleven fragile services, because we're a "
                       "five-person team for the first year.\" If a CEO with a technical background pushes on "
                       "scaling, the honest answer is that a modular monolith splits into services later without "
                       "a rewrite — say that only if asked.")


STACK_NOTES = ("Don't read the table row by row — the room doesn't need to know what Typesense is. The only two "
               "numbers worth saying out loud are the Algolia cost comparison and the AWS Mumbai/DPDP compliance "
               "line, because those are the ones a CEO will actually remember and repeat to an investor.")


TRUST_PIPELINE_NOTES = ("This is the slide to slow down on — say the Puttaswamy line exactly as written on the "
                         "box, word for word, because it's the one legal fact that must not be garbled on stage. "
                         "If asked \"so we can't use Aadhaar at all,\" the answer is we can, just not by calling "
                         "the eKYC API directly — offline XML, DigiLocker, and licensed partners are all legal "
                         "and all in this pipeline.")


BUILD_PHASES_NOTES = ("Anchor the room on the one-city, four-month MVP commitment — that's the number a CEO will "
                       "hold you to, so say \"4 months to a working one-city product\" explicitly rather than "
                       "letting the phase labels speak for themselves.")


TEAM_PLAN_NOTES = ("The headline is \"small team, sequenced hiring\" — resist the urge to justify each hire "
                    "individually. If asked why the founding engineer comes before a designer, the honest answer "
                    "is that the demo app already exists as proof the frontend/product work doesn't block on "
                    "that hire yet.")


BUDGET_NOTES = ("Present all three scenarios before anyone asks which one to pick — this is a menu, not a "
                 "recommendation, and the ask on the next slide is a decision between them, not an approval of a "
                 "single number. If pressed for a personal recommendation, Base is the defensible middle: it "
                 "funds the full hiring sequence without the Lean scenario's slack risk.")


ASK_CLOSE_NOTES = ("This is the actual decision slide — say the ask in one sentence before showing the success "
                    "metrics: \"we need a budget scenario approved and a start date for the founding engineer "
                    "hire.\" The month-12 metrics exist to answer \"how will we know this worked,\" not to be "
                    "read verbatim.")


def add_ask_close_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'THE ASK', 'What we need from you, next 90 days.')
    add_textbox(slide, 640080, 2103120, 10911840, 640080,
                'The decision: approve one budget scenario (Lean / Base / Comfortable) and a start date for the '
                'founding engineer hire.',
                size_pt=15, bold=True, color='graphite', anchor='top')
    add_card(slide, 640080, 2960000, 10911840, 1600000, fill='sand', border='sand_line', adj_raw=3000)
    add_textbox(slide, 868680, 3050000, 10454640, 365760,
                'Next 90 days', size_pt=15, bold=True, color='teal')
    add_textbox(slide, 868680, 3450000, 10454640, 1000000,
                'Phase 0 setup (weeks 0–4) → founding engineer hired and Phase 1 MVP build under way, '
                'targeting the one-city Bengaluru launch inside 4 months.',
                size_pt=13, color='grey', anchor='top')
    add_card(slide, 640080, 4740000, 10911840, 1750000, fill='white', border='sand_line', adj_raw=3000)
    add_textbox(slide, 868680, 4830000, 10454640, 365760,
                'What success looks like at month 12', size_pt=15, bold=True, color='teal')
    add_textbox(slide, 868680, 5230000, 10454640, 1200000,
                '200+ verified owner listings per launch locality  ·  >8% free-to-paid connect conversion  '
                '·  >40% contact-to-visit rate  ·  >70% of live inventory verified.',
                size_pt=13, color='grey', anchor='top')
    set_notes(slide, ASK_CLOSE_NOTES)
    return slide


def add_architecture_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'ARCHITECTURE', 'One well-organised building, not eleven services.')
    add_card(slide, 640080, 2103120, 10911840, 3474720, fill='sand', border='sand_line', adj_raw=3000)
    add_textbox(slide, 868680, 2377440, 10454640, 365760,
                'PropWeb — modular monolith', size_pt=15, bold=True, color='graphite')
    blocks = ['Listings & Search', 'Trust & KYC', 'Matching Engine', 'Connect & Payments']
    n = len(blocks)
    gap = 182880
    inner_w = 10454640
    card_w = (inner_w - gap * (n - 1)) // n
    for idx, label in enumerate(blocks):
        x = 868680 + idx * (card_w + gap)
        add_card(slide, x, 2960000, card_w, 1000000, fill='white', border='sand_line', adj_raw=6061)
        add_textbox(slide, x, 2960000, card_w, 1000000, label, size_pt=13, bold=True,
                    color='teal', align='ctr')
    add_textbox(slide, 640080, 5738160, 10911840, 365760,
                'Small team, one codebase, one language (TypeScript) — ship fast without distributed-systems overhead.',
                size_pt=14, color='grey')
    set_notes(slide, ARCHITECTURE_NOTES)
    return slide


def add_market_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'THE OPPORTUNITY', 'A large market with an unsolved trust gap.')
    add_stat_row(slide, [
        {'number': '$1.3–1.7B', 'label': 'India proptech market size (2025)'},
        {'number': '12–19%', 'label': 'CAGR — proptech growth rate'},
        {'number': '$20B', 'label': 'Indian rental market size'},
        {'number': 'Bengaluru', 'label': 'launch wedge: highest rental velocity, highest brokerage pain, most digitally-ready renters'},
    ])
    add_mixed_textbox(slide, 640080, 4892040, 10972800, 914400, [
        {'text': 'NoBroker already proved renters will pay to skip brokers — ', 'color': 'grey'},
        {'text': '₹803 Cr revenue, 30M+ users', 'color': 'graphite', 'bold': True},
        {'text': ' — inside a market growing double-digit. Demand is proven; the trust gap PropWeb targets is still wide open.', 'color': 'grey'},
    ], size_pt=14, anchor='ctr')
    set_notes(slide, MARKET_NOTES)
    return slide


def add_stack_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'THE STACK', 'A modern stack, corrected for India economics.')
    rows = [
        ("Layer", "Choice", "Why it wins"),
        ("Frontend", "Next.js (React)", "SEO locality pages = free Google traffic"),
        ("Backend", "Node.js + TypeScript", "One language, small team moves fast"),
        ("Database", "PostgreSQL + PostGIS", "'Flats within 3 km' queries, reliably"),
        ("Search", "Typesense (self-hosted)", "~$69/mo vs Algolia ~$525/mo"),
        ("Maps", "Ola Maps / MapmyIndia", "Free tier + best Indian locality data"),
        ("Cloud", "AWS Mumbai", "KYC data stays in India (DPDP compliance)"),
        ("Payments", "Razorpay (UPI-first)", "UPI at 0% MDR by RBI mandate"),
        ("Call masking", "Exotel", "Virtual numbers keep owner phones private"),
        ("KYC APIs", "Surepass/IDfy class vendors", "PAN ≈₹1–3/check; Aadhaar via offline XML/DigiLocker"),
        ("Agreements", "Leegality/Digio", "Aadhaar eSign ₹10–40/signature + state-wise e-stamping"),
    ]
    add_table(slide, 640080, 2103120, 10911840, 3800000, rows,
              col_widths=(2743680, 3200000, 4968160))
    add_textbox(slide, 640080, 6100000, 10911840, 600000,
                "Cost traps avoided: Typesense keeps ~80% of Algolia's power at ~20% of the cost; Google Maps "
                "stays a free-tier fallback, not the primary, ahead of its post-2027 pricing changes.",
                size_pt=12, color='grey', anchor='top')
    set_notes(slide, STACK_NOTES)
    return slide


def add_trust_pipeline_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'THE TRUST PIPELINE', 'Verification, the legal way.')
    add_flow_steps(slide, [
        {'prefix': '01 · IDENTITY', 'title': 'Aadhaar offline XML / DigiLocker',
         'body': 'Identity proof through the legal offline route — no direct Aadhaar eKYC call.'},
        {'prefix': '02 · PAN CHECK', 'title': 'PAN verification API',
         'body': '₹1–3 per check via a licensed KYC API partner.'},
        {'prefix': '03 · OWNERSHIP', 'title': 'Ownership-proof review',
         'body': 'Property tax receipt, electricity bill or sale deed — manual review.'},
        {'prefix': '04 · BADGE', 'title': 'Badge issued',
         'body': 'Verified Owner / Verified Tenant mark goes live on the profile and every listing.'},
    ])
    add_card(slide, 640080, 4892040, 10911840, 1280160, fill='white', border='amber', adj_raw=1500, border_w_emu=19050)
    add_textbox(slide, 868680, 4938840, 10454640, 1188720,
                'The one legal fact that must not be gotten wrong on stage: private companies cannot call Aadhaar '
                'eKYC APIs directly (Supreme Court, Puttaswamy 2018). PropWeb verifies identity only through the '
                'legal routes — Aadhaar offline XML, DigiLocker, or a licensed KYC partner.',
                size_pt=13, color='graphite', anchor='ctr')
    set_notes(slide, TRUST_PIPELINE_NOTES)
    return slide


def add_build_phases_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'ROADMAP', 'Four phases, one city first.')
    add_flow_steps(slide, [
        {'prefix': 'PHASE 0 · weeks 0–4', 'title': 'Setup',
         'body': 'Team hiring starts, tooling and accounts provisioned, design system finalized.'},
        {'prefix': 'PHASE 1 · months 1–4', 'title': 'MVP, one city',
         'body': 'Bengaluru launch: search, listings, trust badges, pay-to-connect live.'},
        {'prefix': 'PHASE 2 · months 4–9', 'title': 'Depth',
         'body': 'AI Match Score, digital rent agreements, fraud/broker detection.'},
        {'prefix': 'PHASE 3 · months 9–15', 'title': 'Services + multi-city',
         'body': 'Affiliate services marketplace launches; expansion beyond Bengaluru.'},
    ])
    set_notes(slide, BUILD_PHASES_NOTES)
    return slide


def add_team_plan_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'TEAM', 'Small team, sequenced hiring.')
    add_flow_steps(slide, [
        {'prefix': 'MONTHS 1–3', 'title': '₹2.75L / month',
         'body': 'Founding full-stack engineer (₹2L/mo) + contract designer (₹75k/mo).'},
        {'prefix': 'MONTHS 4–6', 'title': '₹5.35L / month',
         'body': '+ Backend engineer (₹1.6L/mo) + frontend engineer (₹1L/mo).'},
        {'prefix': 'MONTHS 7–12', 'title': '₹6.65L / month',
         'body': '+ QA/ops (₹70k/mo) + fractional DevOps (₹60k/mo).'},
    ])
    add_textbox(slide, 640080, 5400000, 10911840, 500000,
                '12-month engineering payroll: ₹55–65 lakh.', size_pt=16, bold=True, color='teal')
    set_notes(slide, TEAM_PLAN_NOTES)
    return slide


def add_budget_slide(prs):
    slide = add_blank_slide(prs, light=True)
    add_header(slide, 'BUDGET', 'Three scenarios, months 0–12.')
    rows = [
        ('Cost line', 'Lean', 'Base', 'Comfortable'),
        ('Engineering payroll', '₹40L', '₹60L', '₹90L'),
        ('CTO+CPO cash (reduced)', '₹12L', '₹24L', '₹36L'),
        ('Cloud & infrastructure', '₹3L', '₹6L', '₹12L'),
        ('Third-party APIs', '₹3L', '₹6L', '₹12L'),
        ('Tools & licenses', '₹1.5L', '₹3L', '₹5L'),
        ('Contingency (~15%)', '₹9L', '₹15L', '₹23L'),
        ('TOTAL (0–12 months)', '≈₹68L', '≈₹1.14 Cr', '≈₹1.78 Cr'),
    ]
    add_table(slide, 640080, 2103120, 10911840, 3800000, rows,
              col_widths=(4200000, 2237280, 2237280, 2237280))
    add_textbox(slide, 640080, 6100000, 10911840, 500000,
                'Team of 5–6 + fractional roles; one-city MVP achievable in 4 months.',
                size_pt=13, color='grey')
    set_notes(slide, BUDGET_NOTES)
    return slide


def add_notes_to_existing_slides(prs):
    for marker, note in EXISTING_NOTES.items():
        for slide in prs.slides:
            if marker in get_slide_text(slide):
                set_notes(slide, note)
                break


def build():
    prs = Presentation(SOURCE)
    add_notes_to_existing_slides(prs)
    add_market_slide(prs)
    add_architecture_slide(prs)
    add_stack_slide(prs)
    add_trust_pipeline_slide(prs)
    add_build_phases_slide(prs)
    add_team_plan_slide(prs)
    add_budget_slide(prs)
    add_ask_close_slide(prs)
    # Old combined tech+budget slide is retired only after every new slide has
    # been added: python-pptx's slide-partname generator is a naive
    # len(sldIdLst)+1 counter (pptx/parts/presentation.py _next_slide_partname),
    # so removing a slide and then adding another later reuses a partname that's
    # still in the zip and silently corrupts the deck. Removing last avoids that.
    remove_slide_by_text(prs, 'BUILT TO SCALE')
    reorder_slides(prs, FINAL_ORDER_MARKERS)
    prs.save(OUTPUT)


if __name__ == '__main__':
    build()
